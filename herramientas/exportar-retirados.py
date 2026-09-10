#!/usr/bin/env python3
"""
exportar-retirados.py

Escribe el libro de Excel con lo que se retira del sitio para poder seguir
trabajándolo fuera de línea: los ítems que solo tenían estimación nuestra
(sin una sola cotización real), los proveedores del directorio a los que no
se les pudo confirmar ningún precio, y las categorías que se quedan vacías
al retirar esos ítems.

Se corre ANTES de retirar nada, sobre el estado actual de los datos, y el
resultado se versiona en herramientas/retirados/. Uso:

    python3 herramientas/exportar-retirados.py [salida.xlsx]
"""
import json
import subprocess
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

RAIZ = Path(__file__).resolve().parent.parent
SALIDA = Path(sys.argv[1]) if len(sys.argv) > 1 else RAIZ / "herramientas" / "retirados" / "retirados-del-sitio.xlsx"

FUENTE = "Arial"
AZUL = "1F3864"
TXT = Font(name=FUENTE, size=10)
TIT = Font(name=FUENTE, size=10, bold=True, color="FFFFFF")
H1 = Font(name=FUENTE, size=16, bold=True, color=AZUL)
H2 = Font(name=FUENTE, size=11, bold=True, color=AZUL)
FILL_TIT = PatternFill("solid", fgColor=AZUL)
BORDE = Border(*[Side(style="thin", color="BFBFBF")] * 4)
MONEDA = '#,##0.00;[Red]-#,##0.00;"—"'

NODE = r"""
global.window = global;
const path = require('path');
const D = path.join(process.cwd(), 'precios/assets/js');
['catalogo', 'proveedores', 'precios'].forEach(f => require(path.join(D, 'datos-' + f + '.js')));
const CAT = global.CATALOGO, PROV = global.PROVEEDORES, PRE = global.PRECIOS;
PRE.aplicar(CAT, PROV);
const cuenta = {};
PRE.registros.forEach(q => { cuenta[q.proveedor] = (cuenta[q.proveedor] || 0) + 1; });
const catN = {}; CAT.categorias.forEach(c => { catN[c.codigo] = c; });
const etN = {}; CAT.etapas.forEach(e => { etN[e.codigo] = e.nombre; });
const zoN = {}; (PROV.zonas || []).forEach(z => { zoN[z.codigo] = z.nombre; });
const items = CAT.items.filter(i => i.estado === 'estimado').map(i => ({
  codigo: i.codigo, categoria: i.cat + ' · ' + (catN[i.cat] ? catN[i.cat].nombre : i.cat),
  nombre: i.nombre, esp: i.esp, alcance: i.alcance || '', alias: i.alias || '', unidad: i.unidad,
  etapa: etN[i.etapa] || i.etapa || '', gama: i.gama, origen: i.origen, itbis: i.itbis ? 'Sí' : 'No',
  ref: i.ref, min: i.min, max: i.max, fecha: i.fecha, fuente: i.fuente, nota: i.nota || '',
  medidas: Object.keys(i.medidas || {}).map(k => k + ': ' + i.medidas[k]).join(' · ')
}));
const verificadasPorCat = {}; CAT.items.forEach(i => { if (i.estado !== 'estimado') verificadasPorCat[i.cat] = (verificadasPorCat[i.cat] || 0) + 1; });
const totalPorCat = {}; CAT.items.forEach(i => { totalPorCat[i.cat] = (totalPorCat[i.cat] || 0) + 1; });
const categorias = CAT.categorias.filter(c => !verificadasPorCat[c.codigo]).map(c => ({
  codigo: c.codigo, nombre: c.nombre, desc: c.desc, slug: c.slug, items: totalPorCat[c.codigo] || 0
}));
const proveedores = PROV.lista.filter(p => !p.demo && !cuenta[p.nombre]).map(p => ({
  nombre: p.nombre, tipo: p.tipo, canal: p.canal, publico: p.publico ? 'Sí' : 'No',
  precios: p.precios ? 'Sí' : 'No', cats: (p.cats || []).join(', '),
  catsNombres: (p.cats || []).map(c => catN[c] ? catN[c].nombre : c).join(' · '),
  zonas: (p.zonas || []).map(z => zoN[z] || z).join(', '),
  web: p.web || '', tel: p.tel || '', wa: p.wa || '', email: p.email || '', nota: p.nota || ''
}));
const quedan = PROV.lista.filter(p => !p.demo && cuenta[p.nombre]).map(p => p.nombre + ' (' + cuenta[p.nombre] + ')');
process.stdout.write(JSON.stringify({ items, categorias, proveedores, quedan,
  totales: { items: CAT.items.length, quedanItems: CAT.items.length - items.length, proveedores: PROV.lista.filter(p => !p.demo).length, categorias: CAT.categorias.length },
  generado: new Date().toISOString().slice(0, 10) }));
"""


def datos():
    r = subprocess.run(["node", "-e", NODE], capture_output=True, text=True, cwd=RAIZ)
    if r.returncode != 0:
        sys.stderr.write(r.stderr)
        sys.exit(1)
    return json.loads(r.stdout)


def encabeza(ws, fila, titulos, anchos):
    for i, t in enumerate(titulos, start=1):
        c = ws.cell(row=fila, column=i, value=t)
        c.font, c.fill, c.border = TIT, FILL_TIT, BORDE
        c.alignment = Alignment(vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(i)].width = anchos[i - 1]
    ws.row_dimensions[fila].height = 30
    ws.freeze_panes = ws.cell(row=fila + 1, column=1)


def tabla(ws, fila, filas, monedas=(), ajuste=()):
    for n, f in enumerate(filas, start=fila):
        for i, v in enumerate(f, start=1):
            c = ws.cell(row=n, column=i, value=v)
            c.font = TXT
            c.alignment = Alignment(vertical="top", wrap_text=(i in ajuste))
            if i in monedas:
                c.number_format = MONEDA
    ws.auto_filter.ref = "A%d:%s%d" % (fila - 1, get_column_letter(ws.max_column), max(fila, ws.max_row))


def main():
    d = datos()
    wb = Workbook()
    wb.remove(wb.active)

    ws = wb.create_sheet("Léame")
    ws.column_dimensions["A"].width = 110
    ws["A1"].value = "Lo que se retiró del sitio precios.ingsliberato.com"
    ws["A1"].font = H1
    lineas = [
        "Generado el %s. Nada de esto está publicado: es el material de trabajo para conseguir precios reales." % d["generado"],
        "",
        "Ítems retirados: %d de %d. Son los que solo tenían una estimación nuestra y ninguna cotización real de un comercio." % (len(d["items"]), d["totales"]["items"]),
        "Quedan en el sitio %d ítems, todos con al menos un precio publicado por un comercio (o sin monto por ser de tarifario oficial)." % d["totales"]["quedanItems"],
        "",
        "Categorías retiradas: %d de %d. Se quedaron sin ningún ítem con precio real y por eso salen del sitio con su página." % (len(d["categorias"]), d["totales"]["categorias"]),
        "",
        "Proveedores retirados: %d de %d. Están en el directorio pero no se les pudo confirmar ni un precio." % (len(d["proveedores"]), d["totales"]["proveedores"]),
        "Quedan los que cotizan: " + "; ".join(d["quedan"]) + ".",
        "",
        "Cada ítem conserva su código: al conseguirle un precio real, vuelve al catálogo con ese mismo código.",
        "Los precios de la hoja «Ítems» son la estimación de arranque que tenía el sitio, con su mínimo y máximo: sirven de orden de magnitud, no de referencia.",
    ]
    for i, t in enumerate(lineas, start=3):
        c = ws.cell(row=i, column=1, value=t)
        c.font = TXT
        c.alignment = Alignment(wrap_text=True, vertical="top")

    ws = wb.create_sheet("Ítems")
    cols = [("Código", 12), ("Categoría", 30), ("Ítem", 44), ("Especificación", 40), ("Alcance", 22),
            ("Alias", 24), ("Unidad", 10), ("Etapa", 18), ("Gama", 10), ("Origen", 10), ("ITBIS incl.", 9),
            ("Estimado RD$", 13), ("Mínimo", 12), ("Máximo", 12), ("Fecha", 9), ("Fuente", 34), ("Nota", 34), ("Medidas", 28)]
    encabeza(ws, 1, [c[0] for c in cols], [c[1] for c in cols])
    tabla(ws, 2, [[i["codigo"], i["categoria"], i["nombre"], i["esp"], i["alcance"], i["alias"], i["unidad"], i["etapa"],
                   i["gama"], i["origen"], i["itbis"], i["ref"], i["min"], i["max"], i["fecha"], i["fuente"], i["nota"], i["medidas"]]
                  for i in d["items"]], monedas=(12, 13, 14), ajuste=(3, 4, 16, 17, 18))

    ws = wb.create_sheet("Categorías")
    cols = [("Código", 10), ("Categoría", 34), ("Descripción", 60), ("Página que se retira", 44), ("Ítems que tenía", 12)]
    encabeza(ws, 1, [c[0] for c in cols], [c[1] for c in cols])
    tabla(ws, 2, [[c["codigo"], c["nombre"], c["desc"], c["slug"] + ".html", c["items"]] for c in d["categorias"]], ajuste=(3,))

    ws = wb.create_sheet("Proveedores")
    cols = [("Proveedor", 34), ("Tipo", 14), ("Canal", 12), ("Vende al público", 10), ("Publica precios", 10),
            ("Categorías", 26), ("Categorías (nombre)", 50), ("Zonas", 18), ("Web", 26), ("Teléfono", 14), ("WhatsApp", 14), ("Correo", 28), ("Nota", 60)]
    encabeza(ws, 1, [c[0] for c in cols], [c[1] for c in cols])
    tabla(ws, 2, [[p["nombre"], p["tipo"], p["canal"], p["publico"], p["precios"], p["cats"], p["catsNombres"], p["zonas"],
                   p["web"], p["tel"], p["wa"], p["email"], p["nota"]] for p in d["proveedores"]], ajuste=(7, 13))

    wb.properties.title = "Retirados de precios.ingsliberato.com"
    wb.properties.creator = "Ingenieros Liberato & Asociados"
    SALIDA.parent.mkdir(parents=True, exist_ok=True)
    wb.save(SALIDA)
    print("Escrito %s" % SALIDA.relative_to(RAIZ) if SALIDA.is_relative_to(RAIZ) else SALIDA)
    print("  %d ítems · %d categorías · %d proveedores retirados; quedan %s" % (
        len(d["items"]), len(d["categorias"]), len(d["proveedores"]), ", ".join(d["quedan"])))


if __name__ == "__main__":
    main()
